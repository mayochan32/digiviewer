"""Build the additive Japanese-name snapshot on a maintainer's PC (not at runtime)."""
import argparse
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from datetime import date
import hashlib
import json
from pathlib import Path
import re
import time
import urllib.parse
import urllib.request

ROOT = Path(__file__).resolve().parent.parent
NILIM_URL = 'https://www.nilim.go.jp/lab/fbg/ksnkankyo/mizukokuweb/system/DownLoad/List/R07List/R07zenseibutsu.xlsx'
NILIM_PAGE = 'https://www.nilim.go.jp/lab/fbg/ksnkankyo/mizukokuweb/system/seibutsuListfile.htm'
QUERY = 'SELECT ?taxon ?name ?label WHERE { VALUES ?name { {names} } ?taxon wdt:P225 ?name; wdt:P105 wd:Q7432; rdfs:label ?label . FILTER(LANG(?label) = "ja") }'
SCOPE = 'BioCLIP 2 classification names with positive records in the complete Japan GBIF snapshot; no Japanese Wikipedia article requirement'



def eligible(scientific, japanese):
    # Do not guess a species name from a genus, subspecies, author or partial name.
    return (isinstance(scientific, str) and isinstance(japanese, str)
            and re.fullmatch(r'[A-Z][a-z]+ [a-z][a-z-]+', scientific) is not None
            and scientific.split()[1] not in {'sp', 'spp', 'cf', 'aff'}
            and re.fullmatch(r'[ァ-ヺー・]+', japanese) is not None)


def read_wikidata(data):
    records = []
    for row in data['results']['bindings']:
        scientific, japanese = row['name']['value'], row['label']['value']
        entity = row['taxon']['value'].rsplit('/', 1)[-1]
        if eligible(scientific, japanese) and re.fullmatch(r'Q[1-9][0-9]*', entity):
            records.append((scientific, japanese, 'wikidata', entity, ''))
    return records


def read_nilim(path):
    # This optional build-time dependency is never installed in DigiViewer's environment.
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = book['目録リスト'].iter_rows(values_only=True)
        header = next(rows)
        if any(header[i] != expected for i, expected in {8: '綱', 24: '種コード', 25: '学名', 26: '種和名'}.items()):
            raise ValueError('Unexpected NILIM columns; existing snapshot was not changed')
        records = []
        for row in rows:
            scientific, japanese = row[25:27]
            # Plankton has a separate R08 revision; do not import the obsolete R07-only rows.
            if not any(row[28:34]):
                continue
            if eligible(scientific, japanese) and re.fullmatch(r'[0-9]+', str(row[24])):
                records.append((scientific, japanese, 'nilim-r07', str(row[24]), row[8] or ''))
        return records
    finally:
        book.close()


def merge(base, records):
    grouped = defaultdict(list)
    for record in records:
        grouped[record[0]].append(record)
    entries, preserved, excluded = {}, {}, {}
    for scientific, values in sorted(grouped.items()):
        labels = sorted({row[1] for row in values})
        provenance = sorted({(row[1], row[2], row[3]) for row in values})
        if scientific in base:
            if labels != [base[scientific][0]]:
                preserved[scientific] = {'retained': base[scientific][0], 'sourceNames': provenance}
            continue
        if len(labels) != 1:
            excluded[scientific] = provenance
            continue
        # When labels agree, prefer the domestic list's explicit species record.
        chosen = min(values, key=lambda row: (row[2] != 'nilim-r07', row[3]))
        entries[scientific] = list(chosen[1:4])
    return entries, {'existingNamesPreserved': preserved, 'newNamesExcluded': excluded}


def resolved_keys(entries):
    result = set(entries)
    aliases = json.loads((ROOT / 'src/data/japanese-species-aliases.json').read_text())
    result.update(key for key, value in aliases.items() if value['acceptedName'] in entries)
    return result


def coverage(before, after, classes):
    return {label: {'total': len(names), 'before': len(names & before), 'after': len(names & after)}
            for label, names in sorted(classes.items())}


def model_coverage(labels_path, japan_path, before, after):
    japan = json.loads(japan_path.read_text())
    if not japan.get('complete'):
        raise ValueError('Incomplete Japan occurrence snapshot')
    classes = defaultdict(set)
    for taxonomy, _ in json.loads(labels_path.read_text()):
        # Same name construction as species_worker.py. Denominator is unique names.
        name = ' '.join((taxonomy[5] + ' ' + taxonomy[6]).split()[:2])
        if japan['names'].get(name, 0) > 0:
            classes['all'].add(name)
            if taxonomy[2] in {'Insecta', 'Aves'}:
                classes[taxonomy[2]].add(name)
    return {'japanRetrieved': japan['fetched'], 'labelsSha256': digest(labels_path),
            'japanSha256': digest(japan_path), 'counts': coverage(before, after, classes)}


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def fetch_wikidata(cache, labels_path, japan_path):
    japan = json.loads(japan_path.read_text())
    if not japan.get('complete'):
        raise ValueError('Incomplete Japan occurrence snapshot')
    names = set()
    for taxonomy, _ in json.loads(labels_path.read_text()):
        name = ' '.join((taxonomy[5] + ' ' + taxonomy[6]).split()[:2])
        if japan['names'].get(name, 0) > 0:
            names.add(name)
    names = sorted(names)
    cache.mkdir(parents=True, exist_ok=True)
    def fetch(offset):
        part = names[offset:offset + 150]
        query = QUERY.replace('{names}', ' '.join(json.dumps(name) for name in part))
        target = cache / (hashlib.sha256(query.encode()).hexdigest() + '.json')
        if target.exists():
            return json.loads(target.read_text())['results']['bindings']
        url = 'https://query.wikidata.org/sparql?' + urllib.parse.urlencode({'query': query, 'format': 'json'})
        for attempt in range(3):
            try:
                request = urllib.request.Request(url, headers={
                    'User-Agent': 'DigiViewer (https://github.com/mayochan32/digiviewer; Japanese species names)',
                    'Accept': 'application/sparql-results+json',
                })
                with urllib.request.urlopen(request, timeout=90) as response:
                    data = json.load(response)
                rows = data['results']['bindings']
                target.write_text(json.dumps(data, ensure_ascii=False) + '\n')
                if offset % 1500 == 0:
                    print('Fetched', offset + len(part), '/', len(names), flush=True)
                time.sleep(.2)
                return rows
            except Exception:
                if attempt == 2:
                    raise
                time.sleep(5 * (attempt + 1))
    # Bounded, exact names use the property index and avoid the global query timeout.
    with ThreadPoolExecutor(max_workers=2) as pool:
        rows = [row for part in pool.map(fetch, range(0, len(names), 150)) for row in part]
    return {'queryTemplate': QUERY,
            'queryScope': {'description': SCOPE, 'count': len(names),
                           'namesSha256': hashlib.sha256(json.dumps(names).encode()).hexdigest()},
            'results': {'bindings': rows}}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--wikidata-input', type=Path, help='Complete response including queryScope from the expanded query')
    parser.add_argument('--cache', type=Path, help='Fresh directory for the per-batch query responses')
    parser.add_argument('--nilim-input', type=Path, required=True, help='R07zenseibutsu.xlsx from NILIM_URL')
    parser.add_argument('--retrieved', default=date.today().isoformat(), help='Actual source retrieval date (YYYY-MM-DD)')
    parser.add_argument('--model-labels', type=Path)
    parser.add_argument('--japan-names', type=Path)
    args = parser.parse_args()
    date.fromisoformat(args.retrieved)
    if not args.wikidata_input and not args.cache:
        parser.error('Provide --wikidata-input or --cache')
    if bool(args.model_labels) != bool(args.japan_names):
        parser.error('Provide both --model-labels and --japan-names')
    if not args.wikidata_input and not args.model_labels:
        parser.error('--cache requires --model-labels and --japan-names')
    raw = json.loads(args.wikidata_input.read_text()) if args.wikidata_input else fetch_wikidata(args.cache, args.model_labels, args.japan_names)
    if raw.get('queryTemplate') != QUERY or raw.get('queryScope', {}).get('count', 0) < 10000:
        raise ValueError('Missing or unexpected complete-query provenance')
    wikidata = read_wikidata(raw)
    nilim = read_nilim(args.nilim_input)
    if len({row[0] for row in wikidata}) < 1000 or len({row[0] for row in nilim}) < 28000:
        raise ValueError('Unexpectedly small source; existing snapshot was not changed')
    base_path = ROOT / 'src/data/japanese-species-names.json'
    base = json.loads(base_path.read_text())['entries']
    aliases = json.loads((ROOT / 'src/data/japanese-species-aliases.json').read_text())
    existing = {**base, **{key: base[value['acceptedName']] for key, value in aliases.items()
                         if key not in base and value['acceptedName'] in base}}
    entries, conflicts = merge(existing, wikidata + nilim)
    before, after = resolved_keys(base), resolved_keys({**base, **entries})
    classes = defaultdict(set)
    for scientific, _, _, _, group in nilim:
        if group in {'Insecta', 'Aves'}:
            classes[group].add(scientific)
    report = {
        'retrieved': args.retrieved, 'baseCount': len(base), 'additionalCount': len(entries),
        'combinedCount': len(base) + len(entries),
        'addedBySource': dict(Counter(row[1] for row in entries.values())),
        'wikidataScope': raw['queryScope'],
        'eligibleSourceNames': {'wikidata': len({row[0] for row in wikidata}), 'nilim-r07': len({row[0] for row in nilim})},
        'nilimCoverage': coverage(before, after, classes), **conflicts,
    }
    if args.model_labels:
        report['modelJapanCoverage'] = model_coverage(args.model_labels, args.japan_names, before, after)
    result = {
        'schema': 1, 'retrieved': args.retrieved, 'count': len(entries), 'baseSha256': digest(base_path),
        'sources': {
            'wikidata': {'title': 'Wikidata', 'license': 'CC0-1.0',
                         'licenseUrl': 'https://www.wikidata.org/wiki/Wikidata:Licensing',
                         'retrieved': args.retrieved, 'queryTemplate': QUERY, 'queryScope': raw['queryScope']},
            'nilim-r07': {'title': '国土技術政策総合研究所 河川水辺の国勢調査のための生物リスト（令和7年度）',
                          'license': 'PDL-1.0', 'licenseUrl': 'https://www.nilim.go.jp/aboutlink.htm',
                          'sourceUrl': NILIM_PAGE, 'downloadUrl': NILIM_URL,
                          'retrieved': args.retrieved, 'sha256': digest(args.nilim_input),
                          'recordKey': '目録リスト / 種コード', 'scope': '魚類・底生動物・植物・鳥類・両爬哺・陸上昆虫類等。プランクトン専用行は除外。',
                          'attribution': '国土技術政策総合研究所「河川水辺の国勢調査のための生物リスト（令和7年度）」をDigiViewerが加工して作成。'},
        },
        'entries': entries,
    }
    if args.wikidata_input:
        result['sources']['wikidata']['sha256'] = digest(args.wikidata_input)
    # All parsing and validation finishes before either committed file is changed.
    output = ROOT / 'src/data/japanese-species-names-supplement.json'
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    (ROOT / 'docs/japanese-names-expansion-audit.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({key: value for key, value in report.items() if key not in conflicts}, ensure_ascii=False, indent=2))
    print('Preserved existing disagreements:', len(conflicts['existingNamesPreserved']), 'Excluded new conflicts:', len(conflicts['newNamesExcluded']))


if __name__ == '__main__':
    main()
